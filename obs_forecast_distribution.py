import numpy as np
import polars as pl
from collections import defaultdict
from datetime import date
from argparse import ArgumentParser as ap
from global_functions import date_2_index, city_to_id, get_list_of_weekends, index_2_date
import db_read as dbr
import config_loader as cfg
import yaml
import os
import re
import matplotlib.pyplot as plt
from matplotlib.ticker import FixedLocator
from scipy.stats import linregress
from decimal import Decimal, getcontext, ROUND_HALF_UP, ROUND_HALF_EVEN, InvalidOperation, localcontext
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import time

# ------------------- Globale Konstanten ------------------- #
DEC_QUANT = Decimal('0.001')

# ------------------- Hilfsfunktionen ------------------- #
def to_value(raw_value):
    if raw_value is None:
        return None
    with localcontext() as ctx:
        ctx.prec = 12
        ctx.rounding = ROUND_HALF_EVEN
        d = Decimal(str(raw_value)) / Decimal('10')
        d = d.quantize(DEC_QUANT, rounding=ROUND_HALF_EVEN)
    return float(d)

def get_interval(value, ranges):
    for i, r in enumerate(ranges):
        a, b = r
        if a <= value <= b:
            return i, f"[{a}, {b}]"
    return None, None

def calc_class_means(intervals):
    means = []
    for r in intervals:
        with localcontext() as ctx:
            ctx.prec = 12
            ctx.rounding = ROUND_HALF_EVEN
            a, b = Decimal(str(r[0])), Decimal(str(r[1]))
            means.append((a + b) / 2)
    return means

def _load_yaml_data(filepaths=['tabelle_obs_for.yml']):
    config_data = {}
    for filepath in filepaths:
        with open(filepath, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)
        if data:
            config_data.update(data)
    return config_data

# ------------------- Daten aus DB holen ------------------- #
def get_obs_data(db, staedte, tage, elemente):
    obs_data = {}
    table_name = "wp_wetterturnier_obs"
    cursor = db.con.cursor(dictionary=True)
    for stadt in staedte:
        stationen = cfg.stationen[stadt]
        query = f"""
            SELECT station, betdate, p.paramName, value
            FROM {table_name} w
            INNER JOIN wp_wetterturnier_param p ON w.paramID = p.paramID
            WHERE station IN ({','.join(map(str, stationen))})
              AND betdate IN ({','.join(map(str, tage))})
              AND w.paramID IN ({','.join(map(str, elemente))})
              GROUP BY betdate, p.paramName, station
              ORDER BY betdate ASC, p.sort ASC, station ASC
        """
        cursor.execute(query)
        results = cursor.fetchall()
        nested = {}
        for row in results:
            betdate, param, raw_value = row['betdate'], row['paramName'], row['value']
            if raw_value is not None:
                nested.setdefault(betdate, {}).setdefault(param, []).append(to_value(raw_value))
        obs_data[cfg.id_zu_kuerzel[stadt]] = nested
    return obs_data

def get_forecast_data(db, staedte, tage, elemente, users):
    forecast_data = {}
    table_name = "wp_wetterturnier_bets"
    cursor = db.con.cursor(dictionary=True)
    for stadt in staedte:
        query = f"""
            SELECT betdate, p.paramName, u.user_login, value
            FROM {table_name} w
            INNER JOIN wp_users u ON w.userID = u.ID
            INNER JOIN wp_wetterturnier_param p ON w.paramID = p.paramID
            WHERE w.cityID = {stadt}
              AND betdate IN ({','.join(map(str, tage))})
              AND w.paramID IN ({','.join(map(str, elemente))})
              AND w.userID IN ({','.join(map(str, users))})
            GROUP BY betdate, user_login, p.paramName
            ORDER BY betdate ASC, user_login ASC, p.sort ASC
        """
        cursor.execute(query)
        results = cursor.fetchall()
        nested = {}
        for row in results:
            betdate, user, param, raw_value = row['betdate'], row['user_login'], row['paramName'], row['value']
            if raw_value is not None:
                nested.setdefault(betdate, {}).setdefault(user, {})[param] = to_value(raw_value)
        forecast_data[cfg.id_zu_kuerzel[stadt]] = nested
    return forecast_data

# ------------------- Kombinieren von Obs & Forecast ------------------- #
def combine_data(obs_data, forecast_data, users, users_dict_swapped, elemente_namen):
    combined_data = {}
    for city in obs_data:
        combined_data[city] = {}
        for betdate in obs_data[city]:
            combined_data[city][betdate] = {'o': obs_data[city][betdate], 'f': {}}
            for user in users:
                user_login = users_dict_swapped[user]
                user_login_actual = cfg.teilnehmerumbenennung.get(user_login, user_login)
                try:
                    combined_data[city][betdate]['f'][user_login] = forecast_data[city][betdate].get(
                        user_login_actual,
                        {el: None for el in elemente_namen}
                    )
                except KeyError:
                    combined_data[city][betdate]['f'][user_login] = {el: None for el in elemente_namen}

    return combined_data

# ------------------- Counts & Values by Bin ------------------- #
def calculate_counts(combined_data, elemente_namen, intervals_cfg, users):
    counts = defaultdict(int)
    values_by_bin = defaultdict(list)

    for param in elemente_namen:
        obs_ranges_def = intervals_cfg.get(param, [])
        for_ranges_def = intervals_cfg.get(param, [])
        if not obs_ranges_def or not for_ranges_def:
            print(f"Skipping {param} due to missing ranges.")
            continue

        for city, city_data in combined_data.items():
            for betdate, data in city_data.items():
                obs_vals_list = data['o'].get(param, [])
                valid_obs = [v for v in obs_vals_list if v is not None]
                if not valid_obs:
                    continue
                obs_max = max(valid_obs)
                obs_idx, _ = get_interval(obs_max, obs_ranges_def)
                if obs_idx is None: continue
                obs_range_key = tuple(obs_ranges_def[obs_idx])

                for user, fvals in data['f'].items():
                    fcast_val = fvals.get(param)
                    if fcast_val is None: continue
                    f_idx, _ = get_interval(fcast_val, for_ranges_def)
                    if f_idx is None: continue
                    for_range_key = tuple(for_ranges_def[f_idx])
                    counts[(obs_range_key, for_range_key)] += 1
                    values_by_bin[(obs_range_key, for_range_key)].append((obs_max, fcast_val))

    return counts, values_by_bin

# ------------------- Excel Export ------------------- #
def export_to_excel(combined_data, counts, values_by_bin, elemente_namen, intervals_cfg):
    for param in elemente_namen:
        obs_ranges_def = intervals_cfg.get(param, [])
        for_ranges_def = intervals_cfg.get(param, [])
        if not obs_ranges_def or not for_ranges_def:
            continue

        obs_classes = obs_ranges_def
        fc_classes = for_ranges_def
        n_rows = len(obs_classes)
        n_cols = len(fc_classes)

        all_city_str = "_".join(re.sub(r'[\\/:"*?<>|\s]+', '_', c) for c in combined_data.keys())
        outdir = os.path.join("distribution_outputs", all_city_str)
        os.makedirs(outdir, exist_ok=True)

        users_set = {u for city_data in combined_data.values()
                     for betdate, data in city_data.items()
                     for u in data.get('f', {}).keys()}
        user_str = "_".join(re.sub(r'[\\/:"*?<>|\s]+', '_', u) for u in users_set)
        outfile_xlsx = os.path.join(outdir, f"distribution_all_{all_city_str}_{user_str}.xlsx")

        # Workbook laden oder erstellen
        if os.path.exists(outfile_xlsx):
            wb = load_workbook(outfile_xlsx)
        else:
            wb = Workbook()
            if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1:
                wb.remove(wb["Sheet"])

        # Neues Sheet
        sheet_base_name = f"{param}_{all_city_str}"
        sheet_name = sheet_base_name
        counter = 1
        while sheet_name in wb.sheetnames:
            sheet_name = f"{sheet_base_name}_{counter}"
            counter += 1
        ws = wb.create_sheet(title=sheet_name)

        # Styles
        blue_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
        orchid_fill = PatternFill(start_color="DA70D6", end_color="DA70D6", fill_type="solid")

        # ------------------- Kopfzeilen ------------------- #
        ws.cell(row=1, column=1, value="Kl")

        # ------------------- Matrix Counts ------------------- #
        matrix_counts = [[counts.get((tuple(obs_classes[i]), tuple(fc_classes[j])), 0)
                          for j in range(n_cols)] for i in range(n_rows)]

        # Matrix in Excel schreiben
        for i in range(n_rows):
            for j in range(n_cols):
                cell = ws.cell(row=i+2, column=j+2, value=matrix_counts[i][j])
                if i == j:
                    cell.fill = blue_fill

        # ------------------- Row- und Col-Summen ------------------- #
        row_sums = [sum(row) for row in matrix_counts]
        col_sums = [sum(matrix_counts[i][j] for i in range(n_rows)) for j in range(n_cols)]
        for i, s in enumerate(row_sums):
            ws.cell(row=i+2, column=n_cols+2, value=s)
        for j, s in enumerate(col_sums):
            ws.cell(row=n_rows+2, column=j+2, value=s)
        ws.cell(row=n_rows+2, column=n_cols+2, value=sum(row_sums)).fill = orchid_fill

        # ------------------- Mittelwerte ------------------- #
        
        print(sum(1 for key, pairs in values_by_bin.items()
                     if tuple(key[0]) == tuple(obs_classes[i])
                     for o, f in pairs if o is not None))
        
        
        mob_list = [
            (
                (sum(Decimal(str(o)) for key, pairs in values_by_bin.items()
                     if tuple(key[0]) == tuple(obs_classes[i])  # nur Observation-Klasse
                     for o, f in pairs if o is not None)
                 /
                 sum(1 for key, pairs in values_by_bin.items()
                     if tuple(key[0]) == tuple(obs_classes[i])
                     for o, f in pairs if o is not None)
                ).quantize(Decimal('0.01'))
            ) if any(o is not None for key, pairs in values_by_bin.items()
                     if tuple(key[0]) == tuple(obs_classes[i])
                     for o, f in pairs) else 'NIL'
            for i in range(n_rows)
        ]
        
        mob_sum = sum(float(m) for m in mob_list if m != 'NIL')
        total_sum = sum(col_sums)  # Summe über alle Spalten

        print("Summe mob_list:", mob_sum)
        print("Gesamtsumme col_sums:", total_sum)
        print("Differenz:", total_sum - mob_sum)



     

        mfc_list = [
            (
                (sum(Decimal(str(f)) for key, pairs in values_by_bin.items()
                     if tuple(key[1]) == tuple(fc_classes[j])  # nur Forecast-Klasse
                     for o, f in pairs if f is not None)
                 /
                 sum(1 for key, pairs in values_by_bin.items()
                     if tuple(key[1]) == tuple(fc_classes[j])
                     for o, f in pairs if f is not None)
                ).quantize(Decimal('0.01'))
            ) if any(f is not None for key, pairs in values_by_bin.items()
                     if tuple(key[1]) == tuple(fc_classes[j])
                     for o, f in pairs) else 'NIL'
            for j in range(n_cols)
        ]


        # ------------------- Ergebnisse schreiben ------------------- #
        for i, mfc in enumerate(mfc_list, start=2):
            ws.cell(row=1, column=i, value=mfc)
        for j, mob in enumerate(mob_list, start=2):
            ws.cell(row=j, column=1, value=mob)

        ws.cell(row=n_rows+2, column=1, value="Row_Sum")
        ws.cell(row=1, column=n_cols+2, value="Col_Sum")
        ws.cell(row=n_rows+3, column=1, value="BIAS")

        # ------------------- Bias ------------------- #
        def safe_decimal(val):
            try:
                return Decimal(str(val))
            except (TypeError, ValueError, InvalidOperation):
                return None
        print("Column sums:", col_sums)
        col_bias_list = [
            (
                (
                    sum(
                        (fc - obs) * Decimal(matrix_counts[i][j]) / Decimal(col_sums[j])
                        for i in range(n_rows)
                        if (obs := safe_decimal(ws.cell(row=i+2, column=1).value)) is not None
                        and (fc := safe_decimal(ws.cell(row=1, column=j+2).value)) is not None
                    )
                ).quantize(Decimal("0.01"))
                if col_sums[j] > 0 else "NIL"
            )
            for j in range(n_cols)
        ]


        for j, col_bias in enumerate(col_bias_list, start=2):
            ws.cell(row=n_rows+3, column=j, value=str(col_bias) if col_bias != "NIL" else "NIL")


        # Gewichteter Gesamt-Bias
        gesamt_bias_sum, gesamt_anzahl = map(
            sum,
            zip(*(
                (
                    (Decimal(str(ws.cell(row=1, column=j+2).value)) - Decimal(str(ws.cell(row=i+2, column=1).value))) * Decimal(matrix_counts[i][j]),
                    Decimal(matrix_counts[i][j])
                )
                for i in range(n_rows)
                for j in range(n_cols)
                if ws.cell(row=i+2, column=1).value not in (None, 'NIL')
                and ws.cell(row=1, column=j+2).value not in (None, 'NIL')
            ))
        )
        gesamtbias_weighted = (
            (gesamt_bias_sum / gesamt_anzahl).quantize(Decimal("0.01"))
            if gesamt_anzahl > 0 else "NIL"
        )

        valid_col_bias = [b for b in col_bias_list if b != "NIL"]
        gesamtbias_non_weighted = (
            (sum(valid_col_bias) / len(valid_col_bias)).quantize(Decimal("0.01"))
            if valid_col_bias else "NIL"
        )

        ws.cell(row=n_rows+3, column=n_cols+2, value=str(gesamtbias_weighted) if gesamtbias_weighted != "NIL" else "NIL")
        ws.cell(row=n_rows+4, column=n_cols+2, value=str(gesamtbias_non_weighted) if gesamtbias_non_weighted != "NIL" else "NIL")
        print(mfc_list)
        print(mob_list)
        return col_bias_list


        #wb.save(outfile_xlsx)
        #print(f"Excel table saved (sheet updated): {outfile_xlsx}")


# ------------------- ASCII Export ------------------- #
def export_to_ascii(combined_data, values_by_bin, elemente_namen, intervals_cfg, day_name, col_bias_list):
    """
    Exportiert die Verteilungen als ASCII-Dateien.
    col_bias_list: Liste der Spalten-Bias-Werte wie in Excel berechnet (MFc - Obs)
    """

    def safe_mean(values):
        """Exaktes arithmetisches Mittel mit Decimal, gerundet auf 2 Nachkommastellen."""
        if not values:
            return "NIL"
        total = sum(values, Decimal("0"))
        n = Decimal(len(values))
        return (total / n).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    for param in elemente_namen:
        fc_classes = intervals_cfg.get(param, [])

        all_city_str = "_".join(re.sub(r'[\\/:"*?<>|\s]+', '_', c) for c in sorted(combined_data.keys()))
        outdir = os.path.join("distribution_outputs", all_city_str)
        os.makedirs(outdir, exist_ok=True)

        users_set = {u for city_data in combined_data.values()
                     for _, data in city_data.items()
                     for u in data.get('f', {}).keys()}
        user_str = "_".join(re.sub(r'[\\/:"*?<>|\s]+', '_', u) for u in sorted(users_set))

        asc_outfile = os.path.join(
            outdir,
            f"distribution_{all_city_str}_{param}_{user_str}_{day_name}.asc"
        )

        # ASCII-Header
        col_widths_asc = [5, 6, 6, 4]
        headers = ["Kl", "MFc", "MOb", "#"]
        asc_lines = [
            "  ".join(f"{h:>{w}}" for h, w in zip(headers, col_widths_asc)),
            "  ".join("-" * w for w in col_widths_asc)
        ]

        # ------------------- Berechnung pro FC-Klasse ------------------- #
        for j, (fc_lower, fc_upper) in enumerate(fc_classes):
            fc_vals_total = []
            n_total = 0

            for key, pairs in values_by_bin.items():
                if key[1] == (fc_lower, fc_upper):
                    for o, f in pairs:
                        if f is not None:
                            fc_vals_total.append(Decimal(str(f)))
                            n_total += 1

            mean_fc = safe_mean(fc_vals_total)

            # Bias aus Excel-Spalte übernehmen
            bias = col_bias_list[j] if j < len(col_bias_list) else Decimal("0.0")
            mean_obs = mean_fc - bias if mean_fc != "NIL" and bias != "NIL" else "NIL"

            # Formatierung
            kl_str = f"{Decimal(str(fc_upper)).quantize(Decimal('0.1')):>{col_widths_asc[0]}}"
            mf_str = f"{mean_fc:>{col_widths_asc[1]}.2f}" if mean_fc != "NIL" else f"{mean_fc:>{col_widths_asc[1]}}"
            mo_str = f"{mean_obs:>{col_widths_asc[2]}.2f}" if mean_obs != "NIL" else f"{mean_obs:>{col_widths_asc[2]}}"
            n_str = f"{n_total:>{col_widths_asc[3]}}"

            asc_lines.append("  ".join([kl_str, mf_str, mo_str, n_str]))

        # ------------------- Datei schreiben ------------------- #
        with open(asc_outfile, "w", encoding="utf-8") as f:
            f.write("\n".join(asc_lines) + "\n")
        return col_bias_list


        print(f"ASCII file saved: {asc_outfile}")




# ------------------- Plots ------------------- #
# Hier kommen die Plots. Hier habe ich die individuelle Skalierung für jeden Parameter eingefügt unter den vielen if's.
# Dann habe ich noch den dd12 Plot für jede Stadt als Polarkoordinatenplot hinzugefügt mit verschiedene Farben für die
# Obse und Forecasts.



# ------------------- Ausgabeverzeichnis ------------------- #
outdir = 'distribution_outputs'
plot_outdir = os.path.join(outdir, "plots")
os.makedirs(plot_outdir, exist_ok=True)

# ------------------- YAML laden ------------------- #
with open("cfg.yml") as f:
    cfg_yaml = yaml.safe_load(f)

# ------------------- Hilfsfunktion Achsen ------------------- #
axis_cfg = {
    "sd1":   {"lims": (0, 60), "ticks": 10},
    "sd24":  {"lims": (0, 100), "ticks": 20},
    "ff12":  {"lims": (0, 15), "ticks": 3},
    "fx24":  {"lims": (0, 30), "ticks": 5},
    "tmin":  {"lims": (-15, 25), "ticks": 5},
    "tmax":  {"lims": (-10, 40), "ticks": 5},
    "td12":  {"lims": (-15, 25), "ticks": 5},
}

def set_axis(ax, param, obs_vals, fcast_vals):
    param_lower = param.lower()
    if param_lower in ["rr1", "rr24"]:
        ax.set_xscale("symlog", linthresh=0.1)
        ax.set_yscale("symlog", linthresh=0.1)
        max_val = max(obs_vals.max(), fcast_vals.max())
        ax.set_xlim(0, max_val*1.05)
        ax.set_ylim(0, max_val*1.05)
    else:
        ax.set_xscale("linear")
        ax.set_yscale("linear")
        cfg = axis_cfg.get(param_lower)
        if cfg:
            ticks = np.arange(cfg["lims"][0], cfg["lims"][1]+1, cfg["ticks"])
            ax.set_xlim(cfg["lims"])
            ax.set_ylim(cfg["lims"])
            ax.xaxis.set_major_locator(FixedLocator(ticks))
            ax.yaxis.set_major_locator(FixedLocator(ticks))

# ------------------- Scatter- und Windrosenplots ------------------- #
def plot_parameter_scatter(combined_data, param, selected_days, day_name, plot_outdir, cfg_yaml):
    obs_vals, fcast_vals = [], []

    # --- Einheit aus YAML ermitteln ---
    param_lower = param.lower()
    elemente_neu = [e.lower() for e in cfg_yaml['elemente']['elemente_archiv_neu']]
    elemente_units = cfg_yaml['elemente']['elemente_einheiten_neu']
    if param_lower in elemente_neu:
        si_unit = elemente_units[elemente_neu.index(param_lower)]
    else:
        si_unit = ""  # Fallback

    # --- Beobachtungen & Forecasts sammeln ---
    for city_data in combined_data.values():
        for betdate, data in city_data.items():
            if betdate not in selected_days:
                continue
            obs_list = data["o"].get(param, [])
            if not obs_list:
                continue
            obs_max = max(obs_list)
            for fvals in data["f"].values():
                fcast_val = fvals.get(param)
                if fcast_val is not None:
                    obs_vals.append(obs_max)
                    fcast_vals.append(fcast_val)

    if len(obs_vals) < 2:
        print(f"Not enough data for {param} to plot.")
        return

    obs_vals = np.array(obs_vals)
    fcast_vals = np.array(fcast_vals)

    # --- Regression ---
    slope, intercept, r_value, _, _ = linregress(obs_vals, fcast_vals)

    # --- Frequenz pro Punkt ---
    pairs = np.column_stack((obs_vals, fcast_vals))
    uniq_pairs, idx, counts = np.unique(pairs, axis=0, return_inverse=True, return_counts=True)
    freqs = counts[idx]

    # --- Windrosen für dd12 ---
    if param_lower == "dd12":
        obs_dirs_rad = np.deg2rad(obs_vals)
        fcast_dirs_rad = np.deg2rad(fcast_vals)

        fig = plt.figure(figsize=(8, 8))
        ax = fig.add_subplot(111, polar=True)
        n_bins = 12  # 30° Intervalle
        ax.hist(obs_dirs_rad, bins=n_bins, range=(0, 2*np.pi), alpha=0.6, color="blue", label="Obs")
        ax.hist(fcast_dirs_rad, bins=n_bins, range=(0, 2*np.pi), alpha=0.6, color="red", label="Forecast")

        ax.set_theta_zero_location("N")
        ax.set_theta_direction(-1)
        plt.legend()
        plt.title(f"Wind direction distribution of {param}")

        # Gemeinsamer Ordner für alle Städte
        city_names_combined = "_".join(combined_data.keys())
        city_dir = os.path.join(plot_outdir, city_names_combined)
        os.makedirs(city_dir, exist_ok=True)

        for ext in ["png", "svg"]:
            plt.savefig(os.path.join(city_dir, f"windrose_{param}_{day_name}.{ext}"),
                        dpi=300, bbox_inches='tight', pad_inches=0)
        plt.close(fig)
        print(f"Windrosenplot gespeichert für {param}")
        return

    # --- Scatterplot ---
    fig, ax = plt.subplots(figsize=(12, 8))
    set_axis(ax, param, obs_vals, fcast_vals)

    sc = ax.scatter(obs_vals, fcast_vals, c=freqs, s=50, cmap="coolwarm",
                    alpha=0.7, vmin=freqs.min(), vmax=freqs.max(), clip_on=False)

    # Frequenzen als kleine Zahlen **in die Punkte**
    for (x, y, f) in zip(obs_vals, fcast_vals, freqs):
        ax.text(x, y, str(f), fontsize=6, ha='center', va='center', color='black', weight='bold')

    cbar = plt.colorbar(sc, ax=ax)
    cbar.set_label("Frequency (number of points)")
    cbar.set_ticks(np.arange(freqs.min(), freqs.max()+1, max(1, (freqs.max()-freqs.min())//10)))

    lims = ax.get_xlim()
    ax.plot([lims[0], lims[1]], [lims[0], lims[1]], 'k--', label="Obs = Forecast")
    ax.plot([lims[0], lims[1]], [intercept + slope*lims[0], intercept + slope*lims[1]],
            'r-', label=rf"y={slope:.2f}x+{intercept:.2f}, $R^2={r_value**2:.2f}$")

    ax.set_xlabel(f"Observation ({param}) [{si_unit}]")
    ax.set_ylabel(f"Forecast ({param}) [{si_unit}]")
    day_str = day_name if day_name in ["Sat", "Sun"] else "all days"
    ax.set_title(f"Scatterplot {param} for {day_str} and {', '.join(combined_data.keys())}")

    ax.grid(True)
    ax.legend()

    # --- Gemeinsamer Ordner für alle Städte ---
    city_names_combined = "_".join(combined_data.keys())
    city_dir = os.path.join(plot_outdir, city_names_combined)
    os.makedirs(city_dir, exist_ok=True)

    for ext in ["png", "svg"]:
        plt.savefig(os.path.join(city_dir, f"scatter_{param}_{day_name}.{ext}"),
                    dpi=300, bbox_inches='tight', pad_inches=0)

    plt.close(fig)
    print(f"Scatterplot gespeichert für {param}, Punkte: {len(obs_vals)}, Unique: {len(uniq_pairs)}")

# ------------------- Main-Funktion ------------------- #
def main():
    start = time.time()
    
    # --- DB & Argumentparser initialisieren ---
    db = dbr.db()
    ps = ap()
    ps.add_argument("--von", type=str, default=cfg.datum_neue_elemente)
    ps.add_argument("--bis", type=str, default=cfg.endtermin)
    ps.add_argument("-p", "--params", type=str, default=",".join(cfg.elemente_archiv_neu))
    ps.add_argument("-c", "--cities", type=str, default=",".join(cfg.stadtnamen))
    ps.add_argument("-u", "--users", type=str, default=",".join(cfg.auswertungsteilnehmer))
    ps.add_argument("-d", "--days", type=str, default=",".join(cfg.auswertungstage))
    ps.add_argument("-v", "--verbose", action="store_true")
    ps = ps.parse_args()

    # --- Tage auswählen ---
    tdate_von = date_2_index(ps.von)
    tdate_bis = date_2_index(ps.bis)
    wochenendtage = get_list_of_weekends(tdate_von, tdate_bis)

    if ps.days == "Sat":
        selected_days = [d for d in wochenendtage if index_2_date(d).weekday() == 6]
    elif ps.days == "Sun":
        selected_days = [d for d in wochenendtage if index_2_date(d).weekday() == 0]
    else:
        selected_days = wochenendtage

    # --- Parameter & Städte ---
    elemente_namen = [el for el in ps.params.split(",") if el in cfg.elemente_archiv_neu]
    elemente = db.get_param_ids(elemente_namen).values()
    staedte = [city_to_id(city, cfg) for city in ps.cities.split(",")]
    user_logins = ps.users.split(",")
    users_dict = db.get_user_ids(user_logins)
    users_dict_swapped = {v: k for k, v in users_dict.items()}
    users = users_dict.values()

    # --- YAML Intervalle ---
    intervals_cfg = _load_yaml_data(filepaths=['tabelle_obs_for.yml'])['Intervalle']

    # --- Daten holen & kombinieren ---
    obs_data = get_obs_data(db, staedte, selected_days, elemente)
    forecast_data = get_forecast_data(db, staedte, selected_days, elemente, users)
    combined_data = combine_data(obs_data, forecast_data, users, users_dict_swapped, elemente_namen)

    # --- Zählen & Werte nach Bins ---
    counts, values_by_bin = calculate_counts(combined_data, elemente_namen, intervals_cfg, users)

    # --- Export ---
    export_to_excel(combined_data, counts, values_by_bin, elemente_namen, intervals_cfg)
    col_bias_list = export_to_excel(combined_data, counts, values_by_bin, elemente_namen, intervals_cfg)

    export_to_ascii(
        combined_data,
        values_by_bin,
        elemente_namen,
        intervals_cfg,
        day_name=ps.days,
        col_bias_list=col_bias_list
    )

    # --- Plots ---
    for param in elemente_namen:
        plot_parameter_scatter(
            combined_data, param, selected_days, ps.days, plot_outdir, cfg_yaml
        )

    end = time.time()
    print(f"Laufzeit: {end-start:.2f} Sekunden")

if __name__ == "__main__":
    main()






















